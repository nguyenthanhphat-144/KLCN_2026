# -*- coding: utf-8 -*-
"""Báo cáo KH&CN: giao diện, bộ lọc, Excel và PDF trong một tệp."""

import io
import os
from collections import Counter
from datetime import date, datetime, timedelta
from functools import wraps
from urllib.parse import urlencode
from xml.sax.saxutils import escape

from flask import (
    Blueprint,
    abort,
    redirect,
    render_template_string,
    request,
    send_file,
    session,
    url_for,
)

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.graphics.shapes import Drawing, Rect, String


BLUE = "17365D"

STATES = [
    "Chờ xét duyệt",
    "Đang xét duyệt",
    "Đã duyệt",
    "Từ chối",
    "Yêu cầu bổ sung",
]

FILTER_KEYS = [
    "nhom",
    "nam",
    "tuNgay",
    "denNgay",
    "donVi",
    "maLinhVuc",
    "maLoai",
    "trangThai",
    "kyBaoCao",
    "tacGia",
    "chiMuc",
    "capDeTai",
    "ketQuaNghiemThu",
]


def text(value):
    return "" if value is None else str(value)


def group_name(record):
    name = (
        text(record.get("nhomCongTrinh"))
        + " "
        + text(record.get("tenLoai"))
    ).casefold()

    if "bài báo" in name:
        return "Bài báo"
    if "đề tài" in name:
        return "Đề tài"
    return "Khác"


def distribution(records, key):
    counts = Counter(
        text(record.get(key)).strip() or "Chưa kê khai"
        for record in records
    )
    return sorted(
        counts.items(),
        key=lambda item: (-item[1], item[0]),
    )


def grouped_statistics(records, key):
    result = []

    for label, total in distribution(records, key):
        approved = sum(
            record.get("trangThai") == "Đã duyệt"
            for record in records
            if (
                text(record.get(key)).strip()
                or "Chưa kê khai"
            ) == label
        )

        result.append([
            label,
            total,
            approved,
            approved / total if total else 0,
        ])

    return result


def report_statistics(records):
    records = [dict(record) for record in records]

    articles = [
        record for record in records
        if group_name(record) == "Bài báo"
    ]

    projects = [
        record for record in records
        if group_name(record) == "Đề tài"
    ]

    states = Counter(
        record.get("trangThai")
        for record in records
    )

    total = len(records)
    approved = states["Đã duyệt"]

    funding = 0.0
    missing_funding = 0
    zero_funding = 0

    for record in projects:
        value = record.get("kinhPhi")

        if value is None or value == "":
            missing_funding += 1
            continue

        try:
            amount = float(value)
        except (TypeError, ValueError):
            missing_funding += 1
            continue

        funding += amount
        if amount == 0:
            zero_funding += 1

    summary = [
        ["Tổng hồ sơ", total],
        ["Bài báo", len(articles)],
        ["Đề tài", len(projects)],
        ["Công trình khác", total - len(articles) - len(projects)],
        *[[state, states[state]] for state in STATES],
        ["Tỷ lệ đã duyệt", approved / total if total else 0],
    ]

    sections = [
        (
            "Bài báo theo năm công bố",
            distribution(articles, "namCongBo"),
        ),
        (
            "Bài báo theo loại công trình",
            distribution(articles, "tenLoai"),
        ),
        (
            "Bài báo theo chỉ mục",
            distribution(articles, "chiMucTapChi"),
        ),
        (
            "Bài báo theo trường tác giả",
            distribution(articles, "thongTinTacGia"),
        ),
        (
            "Đề tài theo cấp",
            distribution(projects, "capDeTai"),
        ),
        (
            "Đề tài theo chủ nhiệm",
            distribution(projects, "chuNhiemDeTai"),
        ),
        (
            "Đề tài theo thực hiện/nghiệm thu",
            distribution(projects, "ketQuaNghiemThu"),
        ),
    ]

    return {
        "total": total,
        "articles": len(articles),
        "projects": len(projects),
        "approved": approved,
        "approval_rate": approved / total if total else 0,
        "summary": summary,
        "states": [
            [state, states[state]]
            for state in STATES
        ],
        "units": grouped_statistics(records, "donViChuTri"),
        "fields": grouped_statistics(records, "tenLinhVuc"),
        "sections": sections,
        "funding": funding,
        "missing_funding": missing_funding,
        "zero_funding": zero_funding,
    }


# =========================================================
# EXCEL
# =========================================================

def excel_value(value):
    if isinstance(value, str) and value.startswith(
        ("=", "+", "-", "@")
    ):
        return "'" + value
    return value


def excel_table(
    workbook,
    name,
    title,
    description,
    headers,
    rows,
    widths,
    percentage_columns=(),
):
    sheet = workbook.create_sheet(name)
    last_column = get_column_letter(len(headers))

    for row, value in [
        [1, title],
        [2, "Phạm vi: " + description],
        [
            3,
            "Ngày xuất: "
            + datetime.now().strftime("%d/%m/%Y %H:%M"),
        ],
    ]:
        sheet.merge_cells(
            f"A{row}:{last_column}{row}"
        )

        cell = sheet.cell(
            row, 1, excel_value(value)
        )

        cell.font = Font(
            name="Arial",
            size=16 if row == 1 else 10,
            bold=row == 1,
            color=BLUE,
        )

        cell.alignment = Alignment(
            wrap_text=True,
            vertical="center",
        )

        sheet.row_dimensions[row].height = (
            38 if row < 3 else 24
        )

    for column, label in enumerate(headers, 1):
        cell = sheet.cell(5, column, label)
        cell.font = Font(
            name="Arial",
            bold=True,
            color="FFFFFF",
        )
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(
            wrap_text=True,
            vertical="center",
        )

    sheet.row_dimensions[5].height = 36

    for row_number, values in enumerate(rows, 6):
        estimated_lines = 1

        for column, value in enumerate(values, 1):
            cell = sheet.cell(
                row_number,
                column,
                excel_value(value),
            )

            cell.font = Font(
                name="Arial",
                size=11,
                color="243447",
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            if row_number % 2 == 0:
                cell.fill = PatternFill(
                    "solid",
                    fgColor="F3F6FA",
                )

            if column in percentage_columns:
                cell.number_format = "0.0%"
            elif isinstance(value, (int, float)):
                cell.number_format = "#,##0.##"

            capacity = max(
                6, int(widths[column - 1] * 0.8)
            )

            lines = sum(
                max(
                    1,
                    (len(part) + capacity - 1) // capacity,
                )
                for part in text(value).split("\n")
            )

            estimated_lines = max(
                estimated_lines, lines
            )

        sheet.row_dimensions[row_number].height = min(
            409,
            max(28, estimated_lines * 16 + 8),
        )

    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A6"

    sheet.auto_filter.ref = (
        f"A5:{last_column}{max(sheet.max_row, 5)}"
    )

    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"

    sheet.page_setup.paperSize = (
        sheet.PAPERSIZE_A3
        if len(headers) > 7
        else sheet.PAPERSIZE_A4
    )

    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    sheet.print_title_rows = "1:5"
    sheet.print_area = (
        f"A1:{last_column}{max(sheet.max_row, 5)}"
    )

    sheet.oddFooter.center.text = "Trang &P / &N"

    return sheet


def detail_columns(group):
    if group == "Bài báo":
        return [
            ("maHoSo", "Mã hồ sơ"),
            ("tieuDe", "Tên bài báo"),
            ("thongTinTacGia", "Tác giả"),
            ("donViChuTri", "Đơn vị"),
            ("namCongBo", "Năm công bố"),
            ("tenTapChi", "Tạp chí"),
            ("chiSoISSN", "ISSN"),
            ("doi", "DOI"),
            ("tapSo", "Tập/Số"),
            ("trangBaiBao", "Trang"),
            ("chiMucTapChi", "Chỉ mục"),
            ("trangThai", "Trạng thái hợp lệ"),
        ]

    return [
        ("maHoSo", "Mã hồ sơ"),
        ("maDeTai", "Mã đề tài"),
        ("tieuDe", "Tên đề tài"),
        ("capDeTai", "Cấp đề tài"),
        ("chuNhiemDeTai", "Chủ nhiệm"),
        ("thanhVienThamGia", "Thành viên"),
        ("donViChuTri", "Đơn vị chủ trì"),
        ("thoiGianThucHien", "Thời gian"),
        ("kinhPhi", "Kinh phí (đồng)"),
        ("sanPhamDuKien", "Sản phẩm dự kiến"),
        ("ketQuaNghiemThu", "Thực hiện/nghiệm thu"),
        ("trangThai", "Trạng thái hợp lệ"),
    ]


def generate_excel_report(
    records: list,
    filter_desc: str = "Tất cả hồ sơ KH&CN",
) -> io.BytesIO:
    records = [dict(record) for record in records]
    stats = report_statistics(records)

    workbook = Workbook()
    workbook.remove(workbook.active)

    summary_rows = [
        [label, value]
        for label, value in stats["summary"]
    ]

    summary_rows.extend([
        ["", ""],
        ["Tổng kinh phí đề tài đã kê khai (đồng)", stats["funding"]],
        ["Đề tài chưa có kinh phí", stats["missing_funding"]],
        ["Đề tài lưu kinh phí bằng 0", stats["zero_funding"]],
        [
            "Cách tính tỷ lệ",
            "Số đã duyệt / Tổng hồ sơ trong phạm vi lọc.",
        ],
        [
            "Lưu ý",
            "Duyệt tính hợp lệ không đồng nghĩa nghiệm thu đề tài.",
        ],
    ])

    summary = excel_table(
        workbook,
        "TongHop",
        "BÁO CÁO TỔNG HỢP HỒ SƠ KH&CN",
        filter_desc,
        ["Chỉ tiêu", "Giá trị"],
        summary_rows,
        [55, 42],
    )

    for row in range(6, summary.max_row + 1):
        if text(summary.cell(row, 1).value).startswith("Tỷ lệ"):
            summary.cell(row, 2).number_format = "0.0%"

    summary.auto_filter.ref = None

    chart = BarChart()
    chart.title = "Hồ sơ theo trạng thái"
    chart.y_axis.title = "Số hồ sơ"
    chart.add_data(
        Reference(
            summary,
            min_col=2,
            min_row=10,
            max_row=14,
        ),
        titles_from_data=False,
    )
    chart.set_categories(
        Reference(
            summary,
            min_col=1,
            min_row=10,
            max_row=14,
        )
    )
    chart.legend = None
    chart.width = 20
    chart.height = 10
    summary.add_chart(chart, "D5")

    summary.print_area = (
        f"A1:N{max(30, summary.max_row)}"
    )

    statistical_rows = []

    for title, values in [
        ["Theo đơn vị", stats["units"]],
        ["Theo lĩnh vực đang lưu", stats["fields"]],
    ]:
        statistical_rows.append([
            title, None, None, None
        ])

        statistical_rows.extend(values)

        statistical_rows.append([
            "", None, None, None
        ])

    for title, values in stats["sections"]:
        statistical_rows.append([
            title, None, None, None
        ])

        for label, count in values:
            statistical_rows.append([
                label, count, None, None
            ])

        if not values:
            statistical_rows.append([
                "Không có dữ liệu", None, None, None
            ])

        statistical_rows.append([
            "", None, None, None
        ])

    analysis = excel_table(
        workbook,
        "ThongKe",
        "THỐNG KÊ THEO NHÓM",
        filter_desc,
        [
            "Phân nhóm",
            "Số hồ sơ",
            "Đã duyệt",
            "Tỷ lệ duyệt trong nhóm",
        ],
        statistical_rows,
        [58, 18, 18, 26],
        percentage_columns=(4,),
    )

    analysis.auto_filter.ref = None

    list_rows = []

    for index, record in enumerate(records, 1):
        list_rows.append([
            index,
            record.get("maHoSo", ""),
            record.get("tieuDe", ""),
            record.get("thongTinTacGia", ""),
            record.get("donViChuTri", ""),
            group_name(record),
            record.get("trangThai", ""),
        ])

    listing = excel_table(
        workbook,
        "DanhSach",
        "DANH SÁCH HỒ SƠ KH&CN",
        filter_desc,
        [
            "STT",
            "Mã hồ sơ",
            "Tên công trình",
            "Tác giả",
            "Đơn vị",
            "Nhóm",
            "Trạng thái hợp lệ",
        ],
        list_rows,
        [6, 20, 62, 34, 38, 16, 26],
    )

    listing.freeze_panes = "C6"

    for group, sheet_name in [
        ["Bài báo", "BaiBao"],
        ["Đề tài", "DeTai"],
    ]:
        subset = [
            record for record in records
            if group_name(record) == group
        ]

        if not subset:
            continue

        columns = detail_columns(group)
        detail_rows = []

        for record in subset:
            values = []

            for key, _ in columns:
                value = record.get(key)

                if value is None or value == "":
                    value = "Chưa kê khai"

                values.append(value)

            detail_rows.append(values)

        widths = [
            60 if key in ("tieuDe", "sanPhamDuKien")
            else 35 if key in (
                "thongTinTacGia",
                "thanhVienThamGia",
                "donViChuTri",
                "tenTapChi",
            )
            else 23
            for key, _ in columns
        ]

        sheet = excel_table(
            workbook,
            sheet_name,
            "CHI TIẾT " + group.upper(),
            filter_desc,
            [label for _, label in columns],
            detail_rows,
            widths,
        )

        sheet.freeze_panes = "C6"

    workbook.active = 0

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    return output


# =========================================================
# PDF
# =========================================================

def register_fonts():
    base = os.path.dirname(os.path.abspath(__file__))

    candidates = [
        (
            os.path.join(base, "fonts", "DejaVuSans.ttf"),
            os.path.join(base, "fonts", "DejaVuSans-Bold.ttf"),
        ),
        (
            "C:/Windows/Fonts/arial.ttf",
            "C:/Windows/Fonts/arialbd.ttf",
        ),
        (
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        ),
    ]

    for regular, bold in candidates:
        if os.path.isfile(regular):
            pdfmetrics.registerFont(
                TTFont("KHCNRegular", regular)
            )

            pdfmetrics.registerFont(
                TTFont(
                    "KHCNBold",
                    bold if os.path.isfile(bold) else regular,
                )
            )

            return "KHCNRegular", "KHCNBold"

    raise RuntimeError(
        "Không tìm thấy font tiếng Việt. "
        "Thêm DejaVuSans.ttf và DejaVuSans-Bold.ttf "
        "vào thư mục fonts."
    )


def generate_pdf_report(
    records: list,
    filter_desc: str = "Tất cả hồ sơ KH&CN",
) -> io.BytesIO:
    records = [dict(record) for record in records]
    stats = report_statistics(records)

    regular, bold = register_fonts()
    output = io.BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=30,
        rightMargin=30,
        topMargin=28,
        bottomMargin=34,
    )

    body_style = ParagraphStyle(
        "BodyVN",
        fontName=regular,
        fontSize=9,
        leading=12,
    )

    header_style = ParagraphStyle(
        "HeaderVN",
        parent=body_style,
        fontName=bold,
        textColor=colors.white,
    )

    heading_style = ParagraphStyle(
        "HeadingVN",
        parent=body_style,
        fontName=bold,
        fontSize=12,
        leading=17,
        spaceBefore=14,
        spaceAfter=8,
        keepWithNext=True,
        textColor=colors.HexColor("#" + BLUE),
    )

    def paragraph(value, style=body_style):
        content = escape(text(value)).replace(
            "\n", "<br/>"
        )
        return Paragraph(content or "—", style)

    def make_table(headers, rows, ratios):
        data = [[
            paragraph(label, header_style)
            for label in headers
        ]]

        data.extend(
            [paragraph(value) for value in row]
            for row in rows
        )

        table = Table(
            data,
            colWidths=[
                document.width * ratio
                for ratio in ratios
            ],
            repeatRows=1,
            hAlign="LEFT",
        )

        table.setStyle(TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#" + BLUE),
            ),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [
                    colors.white,
                    colors.HexColor("#F3F6FA"),
                ],
            ),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            (
                "LINEBELOW",
                (0, 0),
                (-1, -1),
                0.4,
                colors.HexColor("#DCE3EB"),
            ),
        ]))

        return table

    elements = []

    logo_path = os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        "static",
        "logo.png",
    )

    if os.path.isfile(logo_path):
        from reportlab.lib.utils import ImageReader

        width, height = ImageReader(
            logo_path
        ).getSize()

        scale = min(100 / width, 45 / height)

        elements.append(
            Image(
                logo_path,
                width=width * scale,
                height=height * scale,
                hAlign="LEFT",
            )
        )

    elements.extend([
        paragraph(
            "BÁO CÁO THỐNG KÊ HỒ SƠ KH&CN",
            heading_style,
        ),
        paragraph("Phạm vi: " + filter_desc),
        paragraph(
            "Ngày xuất: "
            + datetime.now().strftime("%d/%m/%Y %H:%M")
        ),
        paragraph("Tổng hợp", heading_style),
    ])

    summary_rows = [
        [
            label,
            f"{value:.1%}"
            if label.startswith("Tỷ lệ")
            else str(value),
        ]
        for label, value in stats["summary"]
    ]

    elements.append(
        make_table(
            ["Chỉ tiêu", "Giá trị"],
            summary_rows,
            [0.75, 0.25],
        )
    )

    elements.extend([
        Spacer(1, 6),
        paragraph(
            "Tỷ lệ đã duyệt tính trên tổng hồ sơ trong "
            "phạm vi lọc. Duyệt tính hợp lệ không đồng "
            "nghĩa nghiệm thu đề tài."
        ),
        paragraph(
            "Phân bố theo trạng thái",
            heading_style,
        ),
    ])

    drawing = Drawing(document.width, 165)

    maximum = max(
        [count for _, count in stats["states"]] + [1]
    )

    for index, (label, count) in enumerate(stats["states"]):
        y = 140 - index * 28
        bar_width = (
            count / maximum * (document.width - 210)
        )

        drawing.add(String(
            0, y, label,
            fontName=regular,
            fontSize=9,
        ))

        drawing.add(Rect(
            140,
            y - 3,
            bar_width,
            14,
            fillColor=colors.HexColor("#" + BLUE),
            strokeColor=None,
        ))

        drawing.add(String(
            146 + bar_width,
            y,
            str(count),
            fontName=regular,
            fontSize=9,
        ))

    elements.append(drawing)

    for title, rows in [
        ["Theo đơn vị", stats["units"]],
        ["Theo lĩnh vực đang lưu", stats["fields"]],
    ]:
        elements.append(
            paragraph(title, heading_style)
        )

        if rows:
            elements.append(
                make_table(
                    [
                        "Phân nhóm",
                        "Tổng",
                        "Đã duyệt",
                        "Tỷ lệ duyệt",
                    ],
                    [
                        [label, total, approved, f"{rate:.1%}"]
                        for label, total, approved, rate in rows
                    ],
                    [0.55, 0.15, 0.15, 0.15],
                )
            )
        else:
            elements.append(
                paragraph("Không có dữ liệu.")
            )

    for title, rows in stats["sections"]:
        if not rows:
            continue

        elements.extend([
            paragraph(title, heading_style),
            make_table(
                ["Phân nhóm", "Số hồ sơ"],
                rows,
                [0.85, 0.15],
            ),
        ])

    elements.extend([
        paragraph(
            "Kinh phí đề tài",
            heading_style,
        ),
        paragraph(
            f"Tổng kinh phí đã kê khai: "
            f"{stats['funding']:,.0f} đồng. "
            f"Chưa kê khai: {stats['missing_funding']} đề tài. "
            f"Lưu bằng 0: {stats['zero_funding']} đề tài."
        ),
    ])

    for group in ["Bài báo", "Đề tài", "Khác"]:
        subset = [
            record for record in records
            if group_name(record) == group
        ]

        if not subset:
            continue

        elements.append(
            paragraph(
                "Danh sách " + group.lower(),
                heading_style,
            )
        )

        rows = []

        for record in subset:
            description = text(record.get("tieuDe"))
            extra = []
            missing = []

            if group != "Khác":
                for key, label in detail_columns(group):
                    if key in (
                        "maHoSo",
                        "tieuDe",
                        "thongTinTacGia",
                        "donViChuTri",
                        "trangThai",
                    ):
                        continue

                    value = record.get(key)

                    if value is None or value == "":
                        missing.append(label)
                    else:
                        extra.append(
                            label + ": " + text(value)
                        )

            if extra:
                description += "\n" + "; ".join(extra)

            if missing:
                description += (
                    "\nChưa kê khai: "
                    + ", ".join(missing)
                    + "."
                )

            rows.append([
                record.get("maHoSo", ""),
                description,
                text(record.get("thongTinTacGia"))
                + "\n"
                + text(record.get("donViChuTri")),
                record.get("trangThai", ""),
            ])

        elements.append(
            make_table(
                [
                    "Mã hồ sơ",
                    "Công trình và thông tin chi tiết",
                    "Tác giả / Đơn vị",
                    "Trạng thái",
                ],
                rows,
                [0.14, 0.48, 0.24, 0.14],
            )
        )

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(regular, 8)
        canvas.drawString(
            30, 17, "Báo cáo quản lý KH&CN"
        )
        canvas.drawRightString(
            landscape(A4)[0] - 30,
            17,
            f"Trang {doc.page}",
        )
        canvas.restoreState()

    document.build(
        elements,
        onFirstPage=footer,
        onLaterPages=footer,
    )

    output.seek(0)
    return output


# =========================================================
# BỘ LỌC DỮ LIỆU
# =========================================================

def positive_integer(value, label):
    try:
        result = int(value)

        if result < 1:
            raise ValueError()

        return result

    except (TypeError, ValueError):
        abort(
            400,
            description=label + " không hợp lệ.",
        )


def filtered_records():
    # Import tại đây để tránh phụ thuộc vòng khi nạp app.py.
    from database import get_db_connection

    filters = {
        key: request.args.get(key, "").strip()
        for key in FILTER_KEYS
    }

    conditions = ["h.isDeleted = 0"]
    parameters = []
    descriptions = []

    group = filters["nhom"]

    if group not in ("", "baibao", "detai", "khac"):
        abort(400, description="Nhóm công trình không hợp lệ.")

    article_condition = (
        "(lct.nhomCongTrinh LIKE '%Bài báo%' "
        "OR lct.tenLoai LIKE '%Bài báo%')"
    )

    project_condition = (
        "(lct.nhomCongTrinh LIKE '%Đề tài%' "
        "OR lct.tenLoai LIKE '%Đề tài%')"
    )

    if group:
        if group == "baibao":
            conditions.append(article_condition)
        elif group == "detai":
            conditions.append(project_condition)
        else:
            conditions.append(
                f"(NOT {article_condition} "
                f"AND NOT {project_condition})"
            )

        descriptions.append(
            "Nhóm: "
            + {
                "baibao": "Bài báo",
                "detai": "Đề tài",
                "khac": "Khác",
            }[group]
        )

    for key, column, label in [
        ["nam", "h.namCongBo", "Năm công bố/kê khai"],
        ["maLinhVuc", "h.maLinhVuc", "Mã lĩnh vực"],
        ["maLoai", "h.maLoai", "Mã loại công trình"],
    ]:
        if filters[key]:
            number = positive_integer(
                filters[key], label
            )

            conditions.append(column + " = ?")
            parameters.append(number)
            descriptions.append(f"{label}: {number}")

    for key, column, label in [
        ["donVi", "h.donViChuTri", "Đơn vị"],
        ["trangThai", "h.trangThai", "Trạng thái"],
        ["kyBaoCao", "h.kyBaoCao", "Kỳ báo cáo"],
        ["capDeTai", "h.capDeTai", "Cấp đề tài"],
        [
            "ketQuaNghiemThu",
            "h.ketQuaNghiemThu",
            "Thực hiện/nghiệm thu",
        ],
    ]:
        if filters[key]:
            conditions.append(column + " = ?")
            parameters.append(filters[key])
            descriptions.append(
                label + ": " + filters[key]
            )

    if filters["tacGia"]:
        conditions.append(
            "(h.thongTinTacGia LIKE ? "
            "OR h.chuNhiemDeTai LIKE ?)"
        )

        parameters.extend([
            "%" + filters["tacGia"] + "%",
            "%" + filters["tacGia"] + "%",
        ])

        descriptions.append(
            "Tác giả/chủ nhiệm: " + filters["tacGia"]
        )

    if filters["chiMuc"]:
        conditions.append("h.chiMucTapChi LIKE ?")

        parameters.append(
            "%" + filters["chiMuc"] + "%"
        )

        descriptions.append(
            "Chỉ mục: " + filters["chiMuc"]
        )

    parsed_dates = {}

    for key in ["tuNgay", "denNgay"]:
        if filters[key]:
            try:
                parsed_dates[key] = date.fromisoformat(
                    filters[key]
                )
            except ValueError:
                abort(
                    400,
                    description="Ngày không hợp lệ.",
                )

    if (
        "tuNgay" in parsed_dates
        and "denNgay" in parsed_dates
        and parsed_dates["tuNgay"] > parsed_dates["denNgay"]
    ):
        abort(
            400,
            description="Ngày bắt đầu phải trước ngày kết thúc.",
        )

    if "tuNgay" in parsed_dates:
        conditions.append("h.ngayNop >= ?")
        parameters.append(
            parsed_dates["tuNgay"].isoformat()
        )
        descriptions.append(
            "Ngày nộp từ: " + filters["tuNgay"]
        )

    if "denNgay" in parsed_dates:
        try:
            next_day = (
                parsed_dates["denNgay"]
                + timedelta(days=1)
            )
        except OverflowError:
            abort(400, description="Ngày kết thúc quá lớn.")

        conditions.append("h.ngayNop < ?")
        parameters.append(next_day.isoformat())
        descriptions.append(
            "Ngày nộp đến: " + filters["denNgay"]
        )

    sql = """
        SELECT
            h.*,
            lv.tenLinhVuc,
            lct.tenLoai,
            lct.nhomCongTrinh,
            tc.tenTapChi,
            tc.chiSoISSN,
            tc.xepHang,
            ai.tenLinhVuc AS tenLinhVucAI
        FROM HoSoKHCN h
        LEFT JOIN LinhVuc lv
            ON h.maLinhVuc = lv.maLinhVuc
        LEFT JOIN LoaiCongTrinh lct
            ON h.maLoai = lct.maLoai
        LEFT JOIN TapChiHopLe tc
            ON h.maTapChi = tc.maTapChi
        LEFT JOIN LinhVuc ai
            ON h.maLinhVucAI = ai.maLinhVuc
        WHERE
    """

    sql += " AND ".join(conditions)
    sql += " ORDER BY h.maHoSo"

    connection = get_db_connection()

    try:
        rows = connection.execute(
            sql, parameters
        ).fetchall()

        return (
            [dict(row) for row in rows],
            "; ".join(descriptions)
            or "Toàn bộ hồ sơ KH&CN",
            filters,
        )

    finally:
        connection.close()


# =========================================================
# GIAO DIỆN BÁO CÁO TRONG CHÍNH TỆP PYTHON
# =========================================================

REPORT_TEMPLATE = """
<!doctype html>
<html lang="vi">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Báo cáo & Thống kê</title>
<style>
*{box-sizing:border-box}
body{margin:0;background:#f2f5f9;color:#243447;
font-family:Arial,sans-serif;font-size:14px}
main{max-width:1400px;margin:auto;padding:24px}
h1,h2,h3{color:#17365d}
.panel{background:white;padding:22px;border:1px solid #dce3eb;
border-radius:10px;margin-bottom:20px}
.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(210px,1fr));
gap:14px}
label{display:block;margin-bottom:6px;font-weight:bold}
input,select{width:100%;padding:10px;border:1px solid #cbd5e1;
border-radius:6px;background:white}
.btn{display:inline-block;padding:10px 16px;background:#17365d;
color:white;text-decoration:none;border:0;border-radius:6px;
cursor:pointer;margin:4px 4px 4px 0}
.secondary{background:#64748b}.green{background:#187447}
.red{background:#a52d38}
.note{color:#64748b;font-size:13px;line-height:1.6}
.kpi{font-size:30px;font-weight:bold;color:#17365d;margin-top:8px}
.scroll{overflow:auto}
table{border-collapse:collapse;width:100%;font-size:14px}
td,th{padding:11px;border-bottom:1px solid #e2e8f0;
text-align:left;vertical-align:top}
th{background:#edf2f7}
tr:nth-child(even){background:#f8fafc}
.bar{height:12px;background:#17365d;border-radius:3px}
.bar-label{margin:12px 0 5px}
.actions{margin-top:16px}
a{color:#17365d}
</style>
</head>
<body>
<main>
<a href="{{ url_for('dashboard') }}">← Về Dashboard</a>
<h1>Báo cáo & Thống kê</h1>
<p class="note">Các bảng, biểu đồ và file xuất dùng cùng bộ lọc.
Hồ sơ đã xóa mềm không được tính.</p>

<form class="panel" method="get">
<div class="grid">
<div>
<label>Nhóm công trình</label>
<select name="nhom">
{% for value,label in [
('','Tất cả'),('baibao','Bài báo'),
('detai','Đề tài'),('khac','Khác')] %}
<option value="{{ value }}"
{% if filters.nhom==value %}selected{% endif %}>
{{ label }}</option>
{% endfor %}
</select>
</div>

<div>
<label>Năm công bố/kê khai</label>
<input type="number" min="1" max="9999"
name="nam" value="{{ filters.nam }}">
</div>

<div><label>Ngày nộp từ</label>
<input type="date" name="tuNgay" value="{{ filters.tuNgay }}">
</div>

<div><label>Ngày nộp đến</label>
<input type="date" name="denNgay" value="{{ filters.denNgay }}">
</div>

<div>
<label>Đơn vị</label>
<select name="donVi">
<option value="">Tất cả</option>
{% for value in units %}
<option {% if filters.donVi==value %}selected{% endif %}>
{{ value }}</option>
{% endfor %}
</select>
</div>

<div>
<label>Lĩnh vực</label>
<select name="maLinhVuc">
<option value="">Tất cả</option>
{% for value in fields %}
<option value="{{ value.maLinhVuc }}"
{% if filters.maLinhVuc==value.maLinhVuc|string %}selected{% endif %}>
{{ value.tenLinhVuc }}</option>
{% endfor %}
</select>
</div>

<div>
<label>Loại công trình</label>
<select name="maLoai">
<option value="">Tất cả</option>
{% for value in types %}
<option value="{{ value.maLoai }}"
{% if filters.maLoai==value.maLoai|string %}selected{% endif %}>
{{ value.tenLoai }}</option>
{% endfor %}
</select>
</div>

<div>
<label>Trạng thái hợp lệ</label>
<select name="trangThai">
<option value="">Tất cả</option>
{% for value in states %}
<option {% if filters.trangThai==value %}selected{% endif %}>
{{ value }}</option>
{% endfor %}
</select>
</div>

{% for key,label in [
('kyBaoCao','Kỳ báo cáo'),
('tacGia','Tác giả/chủ nhiệm'),
('chiMuc','Chỉ mục bài báo'),
('capDeTai','Cấp đề tài'),
('ketQuaNghiemThu','Thực hiện/nghiệm thu đề tài')
] %}
<div>
<label>{{ label }}</label>
<input name="{{ key }}" value="{{ filters[key] }}">
</div>
{% endfor %}
</div>

<p class="note">
Năm lọc dùng trường năm công bố/kê khai hiện có,
không phải năm thực hiện đề tài.
Kỳ, cấp và nghiệm thu lọc theo giá trị chính xác đã lưu.
Tác giả và chỉ mục tìm theo một phần nội dung.
</p>

<button class="btn" type="submit">Áp dụng</button>
<a class="btn secondary" href="{{ url_for('reports.index') }}">
Xóa bộ lọc</a>
</form>

<div class="panel">
<strong>Phạm vi:</strong> {{ description }}
<div class="actions">
<a class="btn green"
href="{{ url_for('reports.download',kind='excel') }}?{{ query }}">
Tải Excel</a>
<a class="btn red"
href="{{ url_for('reports.download',kind='pdf') }}?{{ query }}">
Tải PDF</a>
</div>
</div>

<div class="grid">
{% for label,value in [
('Tổng hồ sơ',stats.total),
('Bài báo',stats.articles),
('Đề tài',stats.projects),
('Đã duyệt',stats.approved)] %}
<div class="panel">
{{ label }}
<div class="kpi">{{ value }}</div>
</div>
{% endfor %}
</div>

<div class="panel">
<h2>Kết quả kiểm tra tính hợp lệ</h2>
<p>Tỷ lệ đã duyệt:
<strong>{{ '%.1f'|format(stats.approval_rate*100) }}%</strong>
trên tổng hồ sơ trong phạm vi lọc.</p>

{% for label,count in stats.states %}
<div class="bar-label">{{ label }}: {{ count }}</div>
<div class="bar"
style="width:{{ count/stats.total*100 if stats.total else 0 }}%">
</div>
{% endfor %}
</div>

{% for title,rows in [
('Theo đơn vị',stats.units),
('Theo lĩnh vực đang lưu',stats.fields)
] %}
<div class="panel">
<h2>{{ title }}</h2>
<div class="scroll">
<table>
<thead><tr>
<th>Phân nhóm</th><th>Tổng</th>
<th>Đã duyệt</th><th>Tỷ lệ duyệt trong nhóm</th>
</tr></thead>
<tbody>
{% for label,total,approved,rate in rows %}
<tr>
<td>{{ label }}</td><td>{{ total }}</td>
<td>{{ approved }}</td>
<td>{{ '%.1f'|format(rate*100) }}%</td>
</tr>
{% else %}
<tr><td colspan="4">Không có dữ liệu.</td></tr>
{% endfor %}
</tbody>
</table>
</div>
</div>
{% endfor %}

<div class="grid">
{% for title,rows in stats.sections %}
<div class="panel">
<h3>{{ title }}</h3>
<table>
{% for label,count in rows %}
<tr><td>{{ label }}</td><td>{{ count }}</td></tr>
{% else %}
<tr><td>Không có dữ liệu.</td></tr>
{% endfor %}
</table>
</div>
{% endfor %}
</div>

<div class="panel">
<h2>Kinh phí đề tài</h2>
<p>Tổng kinh phí đã kê khai:
<strong>{{ '{:,.0f}'.format(stats.funding) }} đồng</strong>.</p>
<p class="note">
{{ stats.missing_funding }} đề tài chưa có kinh phí;
{{ stats.zero_funding }} đề tài lưu bằng 0.
Giá trị 0 được giữ nguyên, không tự coi là thiếu.
</p>
<p class="note">
Trạng thái duyệt tính hợp lệ không đồng nghĩa nghiệm thu.
Thống kê tác giả dùng nguyên trường tác giả đã lưu.
Cơ sở dữ liệu hiện chưa có tiến độ và năm thực hiện riêng.
</p>
</div>

<div class="panel">
<h2>Danh sách hồ sơ</h2>
<div class="scroll">
<table>
<thead>
<tr><th>Mã</th><th>Tên công trình</th>
<th>Tác giả</th><th>Đơn vị</th><th>Trạng thái</th></tr>
</thead>
<tbody>
{% for record in records %}
<tr>
<td>{{ record.maHoSo }}</td>
<td>{{ record.tieuDe }}</td>
<td>{{ record.thongTinTacGia }}</td>
<td>{{ record.donViChuTri }}</td>
<td>{{ record.trangThai }}</td>
</tr>
{% else %}
<tr><td colspan="5">Không có hồ sơ phù hợp.</td></tr>
{% endfor %}
</tbody>
</table>
</div>

<div class="actions">
Trang {{ page }}/{{ pages }}
{% if page>1 %}
<a class="btn secondary"
href="?{{ query }}&page={{ page-1 }}">Trước</a>
{% endif %}
{% if page<pages %}
<a class="btn secondary"
href="?{{ query }}&page={{ page+1 }}">Sau</a>
{% endif %}
</div>
</div>
</main>
</body>
</html>
"""


# =========================================================
# ĐĂNG KÝ ROUTE
# =========================================================

def staff_only(function):
    @wraps(function)
    def wrapped(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))

        if session.get("vai_tro") not in ("Admin", "CanBo"):
            abort(403)

        return function(*args, **kwargs)

    return wrapped


def register_reports(app):
    # Tránh đăng ký trùng khi người dùng gọi lại.
    if "reports" in app.blueprints:
        return

    reports = Blueprint("reports", __name__)

    @reports.route("/bao-cao")
    @staff_only
    def index():
        from database import get_db_connection

        records, description, filters = filtered_records()

        page = positive_integer(
            request.args.get("page", "1"),
            "Trang",
        )

        pages = max(1, (len(records) + 19) // 20)
        page = min(page, pages)

        connection = get_db_connection()

        try:
            fields = [
                dict(row)
                for row in connection.execute("""
                    SELECT maLinhVuc, tenLinhVuc
                    FROM LinhVuc
                    ORDER BY tenLinhVuc
                """).fetchall()
            ]

            types = [
                dict(row)
                for row in connection.execute("""
                    SELECT maLoai, tenLoai
                    FROM LoaiCongTrinh
                    ORDER BY tenLoai
                """).fetchall()
            ]

            units = [
                row[0]
                for row in connection.execute("""
                    SELECT DISTINCT donViChuTri
                    FROM HoSoKHCN
                    WHERE isDeleted = 0
                    ORDER BY donViChuTri
                """).fetchall()
                if row[0]
            ]

        finally:
            connection.close()

        query = urlencode({
            key: value
            for key, value in filters.items()
            if value
        })

        return render_template_string(
            REPORT_TEMPLATE,
            records=records[(page - 1) * 20:page * 20],
            stats=report_statistics(records),
            description=description,
            filters=filters,
            fields=fields,
            types=types,
            units=units,
            states=STATES,
            page=page,
            pages=pages,
            query=query,
        )

    @reports.route("/bao-cao/tai/<kind>")
    @staff_only
    def download(kind):
        if kind not in ("excel", "pdf"):
            abort(404)

        records, description, _ = filtered_records()

        if kind == "excel":
            stream = generate_excel_report(
                records, description
            )

            extension = "xlsx"

            mimetype = (
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )

        else:
            stream = generate_pdf_report(
                records, description
            )

            extension = "pdf"
            mimetype = "application/pdf"

        filename = (
            "BaoCao_KHCN_"
            + datetime.now().strftime("%Y%m%d_%H%M%S")
            + "."
            + extension
        )

        return send_file(
            stream,
            as_attachment=True,
            download_name=filename,
            mimetype=mimetype,
        )

    app.register_blueprint(reports)

    # Giữ các nút xuất cũ của ứng dụng.
    if "export_excel" in app.view_functions:
        app.view_functions["export_excel"] = (
            lambda: download("excel")
        )

    if "export_pdf" in app.view_functions:
        app.view_functions["export_pdf"] = (
            lambda: download("pdf")
        )
